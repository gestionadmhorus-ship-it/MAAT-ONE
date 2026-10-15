from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import unittest
import shutil
import test_ag_export
from src.safe_export import export_plan, SEPARATED, UNIFIED, folder_name
from src.excel_reader import find_project_excels
from openpyxl import load_workbook

class ExportModeTests(unittest.TestCase):
    def test_modes_repeated_names_and_progress(self):
        with TemporaryDirectory() as directory:
            root=Path(directory)
            plan=test_ag_export.AGExportTests().make_plan(root)
            plan['sets'][0]['anomaly_name']='Daño / aislador'
            second=deepcopy(plan['sets'][0]);second['set_id']='SET-0002'
            plan['sets'].append(second)
            updates=[]
            summary,output,trace=export_plan(plan,SEPARATED,lambda *args:updates.append(args))
            folders=list((output/SEPARATED).iterdir())
            self.assertEqual(len(folders),2)
            self.assertEqual(summary['images_copied'],6)
            self.assertTrue(all(len(list(p.glob('*.JPG')))==3 for p in folders))
            self.assertTrue(all('Daño _ aislador' in p.name for p in folders))
            self.assertTrue(any(len(item)==3 and item[1:]==(6,6) for item in updates))
            wb=load_workbook(trace);self.assertEqual(wb['CARPETAS'].max_row,4);wb.close()
            with TemporaryDirectory() as archived:
                for target in (output/SEPARATED, output/'Anomalias.xlsx', trace):
                    shutil.move(str(target), str(Path(archived)/target.name))
                summary,output,_=export_plan(plan,UNIFIED)
            self.assertEqual(len(list((output/'Anomalias').glob('*.JPG'))),3)
            self.assertEqual(summary['images_copied'],3)
            self.assertEqual(len(find_project_excels(root)),1)

    def test_separated_failure_cleanup_and_names(self):
        self.assertEqual(folder_name('CON'),'_CON')
        self.assertNotEqual(folder_name('...'),'')
        with TemporaryDirectory() as directory:
            root=Path(directory);plan=test_ag_export.AGExportTests().make_plan(root)
            before=set(root.iterdir())
            with patch('src.output_copier.copy2',side_effect=OSError('fallo')):
                with self.assertRaises(ValueError):export_plan(plan,SEPARATED)
            self.assertEqual(set(root.iterdir()),before)

if __name__=='__main__':unittest.main()
