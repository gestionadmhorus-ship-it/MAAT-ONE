from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from openpyxl import Workbook, load_workbook
from src.manual_sets import manual_record, link_zoom, record_details
from src.dji_parser import parse_dji_images
from src.zoom_matcher import match_zoom_images, decide_zoom
from src.set_builder import build_sets
from src.output_planner import build_output_plan
from src.safe_export import export_plan, UNIFIED, SEPARATED

class ManualSetsTests(unittest.TestCase):
    def test_link_preserves_excel_reference_and_requires_review(self):
        record = dict(row_number=6, date='02/10/2026', file_reference='103010', original_data={'elemento':'Poste'})
        image = parse_dji_images([dict(name='DJI_20261002102955_0001_Z.JPG',path='a')])[0]
        result = match_zoom_images([record], [image])[0]
        self.assertEqual(result['status'], 'ZOOM_NOT_FOUND')
        link_zoom(result, image)
        self.assertEqual(result['target_time'], '103010')
        self.assertEqual(result['time_difference'], -15)
        self.assertEqual(result['status'], 'ZOOM_PENDING_REVIEW')
        self.assertFalse(build_sets([result], [image]))
        self.assertIn('02/10/2026 · Hora: 10:30:10', record_details(record,result))
        decide_zoom(result, True)
        self.assertEqual(build_sets([result],[image])[0]['missing_channels'], ['W','T'])

    def test_manual_set_exports_both_modes_and_trace(self):
        for mode in (UNIFIED, SEPARATED):
            with self.subTest(mode=mode), TemporaryDirectory() as directory:
                root=Path(directory)
                wb=Workbook(); wb.save(root/'original.xlsx'); wb.close()
                original=(root/'original.xlsx').read_bytes()
                entries=[]
                for channel in 'WZT':
                    path=root/f'DJI_20261002081530_0007_{channel}.JPG'
                    path.write_bytes(channel.encode())
                    entries.append(dict(name=path.name,path=str(path)))
                images=parse_dji_images(entries)
                zoom=next(i for i in images if i['channel']=='Z')
                record=manual_record(zoom, dict(element='Poste',anomaly='Daño',criticality='Alta',line='RE2',stage='1'),1)
                result=match_zoom_images([record],images)[0];result['manual_added']=True
                link_zoom(result,zoom);decide_zoom(result,True)
                sets=build_sets([result],images);sets[0]['set_id']='SET-0001'
                plan=build_output_plan(root,[result],sets)
                plan['sets'][0].update(source_record=record,anomaly_name='Daño')
                summary,_,trace=export_plan(plan,mode)
                self.assertEqual(summary['images_copied'],3)
                wb=load_workbook(root/'Anomalias.xlsx')
                self.assertEqual(wb.active.cell(2,6).value,'081530')
                self.assertEqual(wb.active.cell(2,7).value,'0007')
                self.assertEqual(wb.active.cell(2,5).value,'Daño');wb.close()
                wb=load_workbook(trace)
                self.assertTrue(any('Agregado manualmente' in str(c.value) for row in wb['TRAZABILIDAD'] for c in row));wb.close()
                self.assertEqual((root/'original.xlsx').read_bytes(),original)
                output=root/('Anomalias' if mode==UNIFIED else 'Anomalias Separadas/Daño')
                self.assertEqual(sorted(p.read_bytes() for p in output.glob('*.JPG')),[b'T',b'W',b'Z'])

    def test_invalid_link_is_not_mutated(self):
        result=dict(normalized_date=None,target_time='103010')
        image=dict(recognized=True,channel='Z')
        before=dict(result)
        with self.assertRaises(ValueError):link_zoom(result,image)
        self.assertEqual(result,before)
