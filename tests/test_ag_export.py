from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import unittest
from openpyxl import Workbook, load_workbook
from src.excel_reader import interpret_rows, read_project_excel, find_project_excels
from src.dji_parser import parse_dji_images
from src.zoom_matcher import match_zoom_images, extract_target_time
from src.set_builder import build_sets
from src.output_planner import build_output_plan
from src.safe_export import export_plan
from src.image_inventory import inventory_project_images

HEADER=['Fecha','Hora','Elemento','Anomalía','Recomendación','Criticidad','Nombre de archivo']

class AGExportTests(unittest.TestCase):
    def test_screenshot_layout_and_invalid_rows(self):
        rows=[['Fecha: 09/10/2026'],HEADER,
              ['09/10/2026','otra hora','Poste','','','',121836],
              ['09/10/2026','','Poste','','','',81530.0],
              ['', '', 'Poste','','','',121836],
              ['09/10/2026','','Poste'],
              ['Fecha finalizada:', '09/10/2026', 'Hora finalizada:', '12:58:03'], []]
        records=interpret_rows(rows)
        self.assertEqual([r['row_number'] for r in records],[3,4,5,6])
        self.assertEqual(extract_target_time(records[1]['file_reference'])[0],'081530')
        images=parse_dji_images([{'name':'DJI_20261009121836_0001_Z.JPG','path':'x'}])
        results=match_zoom_images(records,images)
        self.assertEqual(results[0]['status'],'ZOOM_EXACT')
        self.assertEqual(results[2]['status'],'INVALID_REFERENCE')
        self.assertEqual(results[3]['status'],'INVALID_REFERENCE')
        wrong=match_zoom_images([{'date':'10/10/2026','file_reference':121836}],images)
        self.assertEqual(wrong[0]['status'],'ZOOM_NOT_FOUND')
        for value in [240000,126060,-1,12.5,None,'DJI_20261009121836_0001_Z.JPG']:
            self.assertIsNone(extract_target_time(value))

    def make_plan(self, root):
        wb=Workbook(); ws=wb.active
        ws.append(['Fecha: 09/10/2026']); ws.append(HEADER)
        ws.append(['09/10/2026','12:18:36','Poste','','','',121836])
        wb.save(root/'entrada libre.xlsx')
        folder=root/'cualquier carpeta'; folder.mkdir()
        for channel,second in [('Z','36'),('W','37'),('T','38')]:
            (folder/f'DJI_202610091218{second}_0001_{channel}.JPG').write_bytes(b'original')
        records=read_project_excel(root)
        images,_=inventory_project_images(root)
        parsed=parse_dji_images(images)
        matched=match_zoom_images(records,parsed)
        sets=build_sets(matched,parsed)
        self.assertEqual(sets[0]['status'],'SET_COMPLETE')
        sets[0]['set_id']='SET-0001'
        return build_output_plan(root,matched,sets)

    def test_export_reload_and_repeat(self):
        with TemporaryDirectory() as d:
            root=Path(d); plan=self.make_plan(root)
            summary,dest,trace=export_plan(plan)
            self.assertEqual(summary['images_copied'],3)
            self.assertTrue(trace.is_file())
            wb=load_workbook(trace); self.assertEqual(wb.sheetnames, ["RESUMEN", "TRAZABILIDAD", "INCIDENCIAS", "CARPETAS"]); wb.close()
            self.assertEqual(len(find_project_excels(root)),1)
            images,duplicates=inventory_project_images(root)
            self.assertEqual(len(images),3); self.assertFalse(duplicates)
            self.assertEqual(dest, root.resolve())
            with self.assertRaises(ValueError): export_plan(plan)
            self.assertTrue(trace.exists())
            self.assertFalse(list(root.glob(".maat-one*")))
            self.assertFalse(list(root.glob("Resultados MAAT One*")))
            (root/'segundo.xlsx').touch()
            with self.assertRaises(ValueError): export_plan(plan)

    def test_failure_leaves_no_partial_output(self):
        with TemporaryDirectory() as d:
            root=Path(d); plan=self.make_plan(root)
            before=set(root.iterdir())
            with patch('src.safe_export.generate_traceability_excel',side_effect=OSError('fallo simulado')):
                with self.assertRaises(OSError): export_plan(plan)
            self.assertEqual(set(root.iterdir()),before)
            with patch('src.output_copier.copy2',side_effect=OSError('fallo copia')):
                with self.assertRaises(ValueError): export_plan(plan)
            self.assertEqual(set(root.iterdir()),before)
            source=Path(plan['sets'][0]['absolute_paths']['zoom']); source.unlink()
            with self.assertRaises(FileNotFoundError): export_plan(plan)
            self.assertEqual(set(root.iterdir()),before)

if __name__=='__main__': unittest.main()
