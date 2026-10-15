from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import unittest
import shutil
from openpyxl import load_workbook
import test_ag_export
from src.excel_reader import interpret_rows, find_project_excels
from src.safe_export import export_plan, UNIFIED, SEPARATED
from src.anomalies_excel import HEADERS

class AnomaliesExcelTests(unittest.TestCase):
    def test_original_fields_final_sequence_and_both_modes(self):
        with TemporaryDirectory() as directory:
            root=Path(directory);plan=test_ag_export.AGExportTests().make_plan(root)
            rows=[['Fecha: 09/10/2026','','Origen: TABLET1','','Línea: RE2_H20 | Etapa: 1'],
                  test_ag_export.HEADER,
                  ['09/10/2026','08:15:30','Aislador cerámico','Cabezal dañado','Reemplazar','Media',81530]]
            source=interpret_rows(rows)[0]
            self.assertEqual(source['project_metadata'],{'line':'RE2_H20','stage':'1'})
            item=plan['sets'][0];item['source_record']=source
            old=Path(item['absolute_paths']['zoom'])
            chosen=old.with_name('DJI_20261009081531_0013_Z.JPG');chosen.write_bytes(old.read_bytes())
            item['absolute_paths']['zoom']=str(chosen);item['filenames']['zoom']=chosen.name
            for mode in (UNIFIED,SEPARATED):
                _,output,_=export_plan(plan,mode)
                wb=load_workbook(output/'Anomalias.xlsx');ws=wb.active
                self.assertEqual(tuple(c.value for c in ws[1]),HEADERS)
                self.assertEqual(ws.max_row,2)
                self.assertEqual([ws.cell(2,c).value for c in range(2,9)],
                                 ['RE2_H20',1,'Aislador cerámico','Cabezal dañado','081531','0013','Media'])
                self.assertEqual(ws['A2'].value.strftime('%d/%m/%Y'),'09/10/2026')
                self.assertEqual(ws['G2'].data_type,'s')
                self.assertEqual(ws['G2'].number_format,'@')
                self.assertEqual(ws['A1'].fill.fgColor.rgb,'00E2EFDA')
                folder = output / ('Anomalias' if mode == UNIFIED else SEPARATED)
                exported = list(folder.rglob(chosen.name))
                self.assertEqual(len(exported), 1)
                self.assertEqual(exported[0].read_bytes(), chosen.read_bytes())
                self.assertEqual(ws['F2'].value, exported[0].name.split('_')[1][8:])
                self.assertEqual(ws['G2'].value, exported[0].name.split('_')[2])
                self.assertEqual(ws['A2'].value.strftime('%Y%m%d'), exported[0].name.split('_')[1][:8])
                self.assertEqual(source['file_reference'], 81530)
                wb.close()
                self.assertEqual(len(find_project_excels(root)),1)
                self.assertFalse(list(root.glob('.maat-one*')))
                with TemporaryDirectory() as archived:
                    for target in (folder, output/'Anomalias.xlsx', output/f'Trazabilidad MAAT One {root.name}.xlsx'):
                        shutil.move(str(target), str(Path(archived)/target.name))
            self.assertEqual(len(find_project_excels(root)),1)

    def test_second_excel_failure_does_not_publish(self):
        with TemporaryDirectory() as directory:
            root=Path(directory);plan=test_ag_export.AGExportTests().make_plan(root)
            before=set(root.iterdir())
            with patch('src.safe_export.generate_anomalies_excel',side_effect=OSError('fallo')):
                with self.assertRaises(OSError):export_plan(plan)
            self.assertEqual(set(root.iterdir()),before)

if __name__=='__main__':unittest.main()
