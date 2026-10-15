from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import unittest
from PIL import Image
from openpyxl import Workbook, load_workbook
from src.ui.main_window import MainWindow
from src.excel_reader import read_project_excel
from src.dji_parser import parse_dji_images
from src.zoom_matcher import match_zoom_images
from src.output_planner import build_output_plan
from src.safe_export import export_plan, UNIFIED, SEPARATED

class AcceptVisibleTests(unittest.TestCase):
    def test_visible_0083_replaces_0085_and_exports_in_both_modes(self):
        for mode in (UNIFIED, SEPARATED):
            with self.subTest(mode=mode), TemporaryDirectory() as d:
                root=Path(d)
                wb=Workbook();ws=wb.active
                ws.append(['Fecha','Hora','Elemento','Anomalía','Recomendación','Criticidad','Nombre de archivo'])
                ws.append(['02/10/2026','','Atadura','Floja','','Baja','101954'])
                wb.save(root/'entrada.xlsx');wb.close()
                images=[]
                for seq, time in [('0083','101930'),('0085','101952')]:
                    for channel in 'WZT':
                        p=root/f'DJI_20261002{time}_{seq}_{channel}.JPG'
                        Image.new('RGB',(40,40),'red' if seq=='0083' else 'blue').save(p)
                        images.append(dict(name=p.name,path=str(p)))
                app=MainWindow()
                try:
                    app.current_root=root;app.review_active=True
                    app._records=read_project_excel(root);app._images=images
                    app._parsed_images=parse_dji_images(images);app._recognized_count=6;app._unrecognized=[]
                    app.zoom_results=match_zoom_images(app._records,app._parsed_images)
                    app.pending_reviews=list(app.zoom_results)
                    app._refresh_summary();app._refresh_review_list()
                    result=app.zoom_results[0]
                    self.assertEqual(result['zoom']['sequence'],'0085')
                    selected=next(i for i in app._parsed_images if i['sequence']=='0083' and i['channel']=='Z')
                    app._select_alternative(selected)
                    self.assertEqual(app._left_images['Zoom']['sequence'],'0083')
                    self.assertEqual(app._left_images['Wide']['sequence'],'0083')
                    self.assertEqual(result['zoom']['sequence'],'0085')
                    for seconds, color in [(0,'#15803D'),(-5,'#15803D'),(6,'#C06000'),(-10,'#C06000'),(11,'#DC2626')]:
                        self.assertEqual(app._time_color(seconds),color)
                    app._open_image_popup(app._left_images['Wide'])
                    app.root.update()
                    import tkinter as tk
                    popup=next(w for w in app.root.winfo_children() if isinstance(w,tk.Toplevel))
                    canvas=popup.winfo_children()[0]
                    canvas.event_generate('<MouseWheel>',delta=120,x=100,y=100)
                    canvas.event_generate('<ButtonPress-1>',x=100,y=100)
                    canvas.event_generate('<B1-Motion>',x=120,y=130)
                    app.root.update();popup.destroy()
                    self.assertEqual(result['zoom']['sequence'],'0085')
                    self.assertEqual(app._viewed_zoom_candidate['sequence'],'0083')
                    app.output_plan={'stale':True}
                    with patch('src.ui.main_window.messagebox.showinfo'):
                        app._decide_current_zoom(True)
                    self.assertEqual(result['zoom']['sequence'],'0083')
                    self.assertEqual(result['original_suggested_zoom']['sequence'],'0085')
                    self.assertEqual(result['time_difference'],-24)
                    self.assertIsNone(app.output_plan)
                    self.assertTrue(all(app.set_results[0][k]['sequence']=='0083' for k in ('wide','zoom','thermal')))
                    plan=build_output_plan(root,app.zoom_results,app.set_results)
                    plan['sets'][0].update(source_record=app._records[0],anomaly_name='Floja')
                    _,_,trace=export_plan(plan,mode)
                    out=root/('Anomalias' if mode==UNIFIED else 'Anomalias Separadas/Floja')
                    self.assertEqual(len(list(out.glob('*0083*.JPG'))),3)
                    self.assertEqual(len(list(out.glob('*0085*.JPG'))),0)
                    wb=load_workbook(root/'Anomalias.xlsx');self.assertEqual(wb.active.cell(2,6).value,'101930');self.assertEqual(wb.active.cell(2,7).value,'0083');wb.close()
                    wb=load_workbook(trace);rows=list(wb['TRAZABILIDAD'].values);record=dict(zip(rows[0],rows[1]));wb.close()
                    self.assertIn('0083',record['Zoom finalmente utilizada'])
                    self.assertIn('0085',record['Zoom sugerida originalmente'])
                    # Revisitar y aceptar la otra imagen también sustituye la elección anterior.
                    other=next(i for i in app._parsed_images if i['sequence']=='0085' and i['channel']=='Z')
                    app._select_alternative(other)
                    app._decide_current_zoom(True)
                    self.assertEqual(result['zoom']['sequence'],'0085')
                    app._exclude_set()
                    self.assertEqual(len(app.anomaly_rows),0)
                    self.assertEqual(len(app.set_results),0)
                    excluded=build_output_plan(root,app.zoom_results,app.set_results)
                    self.assertEqual(excluded['sets'],[])
                    self.assertEqual(excluded['records'][0]['status'],'EXCLUDED')
                    app._undo_exclude()
                    self.assertEqual(len(app.anomaly_rows),1)
                    self.assertEqual(app.set_results[0]['zoom']['sequence'],'0085')
                finally:
                    app.root.update_idletasks()
                    app.root.destroy()
