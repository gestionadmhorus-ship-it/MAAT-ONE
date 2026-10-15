"""Generación del Excel final de trazabilidad de MAAT One."""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


TRACE_HEADERS = (
    "SET ID", "Fila Excel", "Referencia Excel", "Fecha", "Hora esperada",
    "Estado Zoom", "Validación Zoom", "Aprobada", "Archivo Zoom",
    "Zoom sugerida originalmente", "Zoom finalmente utilizada",
    "Motivo cambio Zoom", "Distancia cambio GPS (m)",
    "Hora Zoom", "Diferencia Zoom (s)", "Secuencia", "Estado SET",
    "Archivo Wide", "Diferencia Wide (s)", "Archivo Thermal",
    "Diferencia Thermal (s)", "Canales faltantes", "Canales ambiguos",
    "Observaciones",
)
INCIDENT_HEADERS = (
    "SET ID", "Fila Excel", "Referencia", "Tipo de incidencia", "Detalle",
    "Archivo relacionado", "Diferencia (s)", "Decisión",
)
THIN_BORDER = Border(
    bottom=Side(style="thin", color="D9DEE5")
)
HEADER_FILL = PatternFill("solid", fgColor="DCE6F1")
COMPLETE_FILL = PatternFill("solid", fgColor="E2F0D9")
INCOMPLETE_FILL = PatternFill("solid", fgColor="FFF2CC")
ERROR_FILL = PatternFill("solid", fgColor="FCE4D6")


class TraceabilityError(Exception):
    """Error controlado al generar la trazabilidad final."""


def _approved_text(value: object) -> str:
    if value is True:
        return "Sí"
    if value is False:
        return "No"
    return "—"


def _channel_names(channels: list[str]) -> str:
    names = {"W": "Wide", "Z": "Zoom", "T": "Thermal"}
    return ", ".join(names.get(channel, channel) for channel in channels)


def _ambiguous_text(channels: dict[str, list[str]]) -> str:
    names = {"W": "Wide", "Z": "Zoom", "T": "Thermal"}
    return "; ".join(
        f"{names.get(channel, channel)}: {' | '.join(files)}"
        for channel, files in channels.items()
    )


def _signed_seconds(value: object) -> str:
    if value is None:
        return ""
    return f"{value:+d} s"


def _observations(record: dict[str, object]) -> str:
    status = record["status"]
    difference = record["time_difference"]
    zoom_replaced = (
        record["original_suggested_filename"]
        != record["final_used_filename"]
    )
    messages = {
        "EXCLUDED": "SET excluido por el usuario; no se exportan sus imágenes ni su fila de anomalías.",
        "ZOOM_EXACT": "Zoom exacta validada automáticamente.",
        "ZOOM_REJECTED": "Zoom rechazada manualmente.",
        "ZOOM_NOT_FOUND": "Zoom no encontrada dentro de ±5 s.",
        "ZOOM_AMBIGUOUS": "Zoom ambigua.",
        "INVALID_REFERENCE": "Fecha en A u hora HHMMSS en G no válida.",
    }
    parts: list[str] = []
    if record.get("manual_added"):
        parts.append("Agregado manualmente; no existe fila en el Excel original.")
    if record.get("override_reason") and not zoom_replaced:
        parts.append(str(record["override_reason"]))
    if status == "ZOOM_ACCEPTED" and not zoom_replaced:
        parts.append(
            "Zoom aceptada manualmente con "
            f"{_signed_seconds(difference)} de diferencia."
        )
    elif status in messages:
        parts.append(messages[status])

    trace_set = record["set"]
    if trace_set is not None:
        if trace_set["status"] == "SET_INCOMPLETE":
            missing = _channel_names(trace_set["missing_channels"])
            parts.append(f"SET incompleto: falta {missing}.")
        else:
            parts.append("SET completo.")
    if zoom_replaced:
        replacement = (
            "Zoom sustituida manualmente. "
            f"Original: {record['original_suggested_filename']}. "
            f"Final: {record['final_used_filename']}. "
            f"Motivo: {record['override_reason'] or 'Sin especificar'}. "
        )
        if record["override_distance_m"] is None:
            replacement += "Distancia GPS no disponible."
        else:
            distance = f"{record['override_distance_m']:.1f}".replace(
                ".", ","
            )
            replacement += f"Distancia estimada: {distance} m."
        parts.append(replacement)
    return " ".join(parts)


def _trace_row(record: dict[str, object]) -> tuple[object, ...]:
    trace_set = record["set"]
    if trace_set is None:
        set_id = sequence = ""
        set_status = "SIN SET"
        wide_name = wide_difference = thermal_name = thermal_difference = ""
        missing = ambiguous = ""
    else:
        set_id = trace_set["set_id"]
        sequence = trace_set["sequence"]
        set_status = trace_set["status"]
        wide_name = trace_set["wide"]["filename"] or ""
        wide_difference = trace_set["wide"]["difference_seconds"]
        thermal_name = trace_set["thermal"]["filename"] or ""
        thermal_difference = trace_set["thermal"]["difference_seconds"]
        missing = _channel_names(trace_set["missing_channels"])
        ambiguous = _ambiguous_text(trace_set["ambiguous_channels"])
    return (
        set_id, record["row_number"], record["file_reference"],
        record["normalized_date"] or "", record["target_time"] or "",
        record["status"], record["validation"] or "", 
        _approved_text(record["approved"]), record["filename"] or "",
        record["original_suggested_filename"] or "",
        record["final_used_filename"] or "",
        record["override_reason"] or "",
        record["override_distance_m"],
        record["zoom_time"] or "", record["time_difference"], sequence,
        set_status, wide_name, wide_difference, thermal_name,
        thermal_difference, missing, ambiguous, _observations(record),
    )


def _incident_row(record: dict[str, object]) -> tuple[object, ...] | None:
    trace_set = record["set"]
    kinds: list[str] = []
    details: list[str] = []
    status = record["status"]
    zoom_replaced = (
        record["original_suggested_filename"]
        != record["final_used_filename"]
    )
    if status in {
        "ZOOM_REJECTED", "ZOOM_NOT_FOUND", "ZOOM_AMBIGUOUS",
        "INVALID_REFERENCE", "EXCLUDED",
    }:
        kinds.append(status)
        details.append(_observations(record))
    if trace_set is not None and trace_set["status"] == "SET_INCOMPLETE":
        kinds.append("SET_INCOMPLETE")
        details.append(
            f"Canales faltantes: {_channel_names(trace_set['missing_channels'])}"
        )
    if trace_set is not None and trace_set["ambiguous_channels"]:
        kinds.append("CANALES_AMBIGUOS")
        details.append(_ambiguous_text(trace_set["ambiguous_channels"]))
    if (
        status == "ZOOM_ACCEPTED"
        and record["time_difference"] != 0
        and not zoom_replaced
    ):
        kinds.append("ZOOM_ACCEPTED_CON_DIFERENCIA")
        details.append(_observations(record))
    if zoom_replaced:
        kinds.append("ZOOM_REEMPLAZADA_MANUALMENTE")
        distance = (
            "Distancia GPS no disponible"
            if record["override_distance_m"] is None
            else "Distancia GPS: "
            + f"{record['override_distance_m']:.1f} m".replace(".", ",")
        )
        details.append(
            f"Original: {record['original_suggested_filename']}; "
            f"Final: {record['final_used_filename']}; "
            f"Motivo: {record['override_reason'] or 'Sin especificar'}; "
            f"{distance}"
        )
    if not kinds:
        return None
    related = (
        record["final_used_filename"]
        if zoom_replaced
        else record["filename"] or " | ".join(record["candidates"])
    )
    return (
        trace_set["set_id"] if trace_set is not None else "",
        record["row_number"], record["file_reference"], "; ".join(kinds),
        "; ".join(dict.fromkeys(details)), related,
        record["time_difference"] if record["time_difference"] is not None else "",
        _approved_text(record["approved"]),
    )


def _style_header(sheet, columns: int) -> None:
    for cell in sheet[1][:columns]:
        cell.font = Font(bold=True, color="1F2937")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    sheet.row_dimensions[1].height = 30


def _build_summary(workbook: Workbook, plan: dict[str, object]) -> None:
    sheet = workbook.active
    sheet.title = "RESUMEN"
    root_name = Path(plan["root_folder"]).name
    summary = plan["summary"]
    sheet["A1"] = "MAAT One"
    sheet["A1"].font = Font(size=22, bold=True, color="1F4E78")
    sheet["A2"] = "Trazabilidad de procesamiento"
    sheet["A2"].font = Font(size=14, bold=True, color="44546A")
    rows = (
        ("Proyecto", root_name),
        ("Carpeta de salida", plan["destination_name"]),
        ("Registros procesados", len(plan["records"])),
        ("SET completos", summary["complete_sets"]),
        ("SET incompletos", summary["incomplete_sets"]),
        ("SET descartados", summary["discarded_sets"]),
        ("Imágenes previstas para copia", summary["images_to_copy"]),
        ("Imágenes faltantes", summary["missing_images"]),
        ("Decisiones manuales", summary["manual_decisions"]),
        ("Validaciones automáticas", summary["automatic_validations"]),
        ("Fecha y hora de generación", datetime.now()),
    )
    for row, (label, value) in enumerate(rows, start=4):
        sheet.cell(row, 1, label).font = Font(bold=True)
        sheet.cell(row, 1).fill = HEADER_FILL
        sheet.cell(row, 2, value)
        for column in (1, 2):
            cell = sheet.cell(row, column)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[row].height = 22
    sheet.cell(14, 2).number_format = "dd/mm/yyyy hh:mm:ss"
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 52


def _build_traceability(workbook: Workbook, plan: dict[str, object]) -> None:
    sheet = workbook.create_sheet("TRAZABILIDAD")
    sheet.append(TRACE_HEADERS)
    for record in plan["records"]:
        sheet.append(_trace_row(record))
        sheet.cell(sheet.max_row, 13).number_format = "0.0"
        fill = None
        trace_set = record["set"]
        if record["status"] in {
            "ZOOM_REJECTED", "ZOOM_NOT_FOUND", "ZOOM_AMBIGUOUS",
            "INVALID_REFERENCE",
        }:
            fill = ERROR_FILL
        elif trace_set is not None and trace_set["status"] == "SET_INCOMPLETE":
            fill = INCOMPLETE_FILL
        elif trace_set is not None and trace_set["status"] == "SET_COMPLETE":
            fill = COMPLETE_FILL
        for cell in sheet[sheet.max_row]:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if fill is not None:
                cell.fill = fill
        sheet.row_dimensions[sheet.max_row].height = 34
    _style_header(sheet, len(TRACE_HEADERS))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = (
        14, 11, 22, 12, 14, 22, 18, 11, 28, 30, 30, 22,
        22, 12, 20, 12, 18, 28, 20, 28, 22, 22, 44, 55,
    )
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width


def _build_incidents(workbook: Workbook, plan: dict[str, object]) -> None:
    sheet = workbook.create_sheet("INCIDENCIAS")
    sheet.append(INCIDENT_HEADERS)
    for record in plan["records"]:
        row = _incident_row(record)
        if row is not None:
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                cell.border = THIN_BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.fill = ERROR_FILL
            sheet.row_dimensions[sheet.max_row].height = 34
    _style_header(sheet, len(INCIDENT_HEADERS))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = (14, 11, 22, 34, 58, 35, 16, 14)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width


def generate_traceability_excel(plan: dict[str, object] | None, output_path: Path | None = None) -> Path:
    """Valida el plan y guarda un libro nuevo sin sobrescribir otro."""
    if plan is None:
        raise TraceabilityError("No existe un plan de salida.")
    if "records" not in plan:
        raise TraceabilityError("El plan no contiene trazabilidad completa.")
    if any(
        record.get("status") == "ZOOM_PENDING_REVIEW"
        for record in plan["records"]
    ):
        raise TraceabilityError("Existen revisiones de Zoom pendientes.")
    destination_folder = Path(plan["destination_folder"])
    if not destination_folder.is_dir():
        raise TraceabilityError("No existe la carpeta de copia.")

    root_folder = Path(plan["root_folder"])
    if output_path is None:
        output_path = root_folder / f"Trazabilidad MAAT One {root_folder.name}.xlsx"
    if output_path.exists():
        raise TraceabilityError(
            "Ya existe el Excel de trazabilidad de MAAT One."
        )

    workbook = Workbook()
    if plan.get("output_identifier"):
        workbook.properties.identifier = plan["output_identifier"]
    _build_summary(workbook, plan)
    _build_traceability(workbook, plan)
    _build_incidents(workbook, plan)
    if plan.get("export_folders"):
        sheet = workbook.create_sheet("CARPETAS")
        sheet.append(["Modo", plan["export_mode"]])
        sheet.append(["SET", "Anomalía", "Carpeta relativa"])
        for entry in plan["export_folders"]: sheet.append(list(entry))
        sheet.column_dimensions["A"].width = 25
        sheet.column_dimensions["B"].width = 50
        sheet.column_dimensions["C"].width = 70
    try:
        workbook.save(output_path)
    except OSError as error:
        raise TraceabilityError(
            "No fue posible guardar el Excel de trazabilidad."
        ) from error
    return output_path
