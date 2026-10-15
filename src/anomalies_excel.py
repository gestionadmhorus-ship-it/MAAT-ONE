"""Segundo Excel: datos de origen y secuencia de la Zoom finalmente exportada."""
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from src.dji_parser import parse_dji_image

HEADERS = ("FECHA", "LINEA", "ETAPA", "ELEMENTO", "ANOMALIA", "HORA DE ARCHIVO", "NUMERO DE ARCHIVO", "CRITICIDAD")

def generate_anomalies_excel(plan, output_path):
    output_path = Path(output_path)
    if output_path.exists(): raise ValueError("Ya existe el Excel de anomalías.")
    workbook = Workbook()
    if plan.get("output_identifier"):
        workbook.properties.identifier = plan["output_identifier"]
    sheet = workbook.active
    sheet.title = "ANOMALIAS"
    sheet.append(HEADERS)
    border = Border(bottom=Side(style="thin", color="D9E1D3"))
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="E2EFDA")
        cell.font = Font(name="Calibri", size=11, bold=True)
        cell.alignment = Alignment(vertical="center")
    for item in plan["sets"]:
        source = item.get("source_record")
        if source is None: raise ValueError("Faltan los datos originales del SET " + item["set_id"])
        zoom = parse_dji_image({"name": item["filenames"]["zoom"], "path": item["absolute_paths"]["zoom"]})
        if not zoom.get("recognized") or zoom.get("channel") != "Z" or Path(item["absolute_paths"]["zoom"]).name != item["filenames"]["zoom"]:
            raise ValueError("Fecha, hora o Zoom inválida para " + item["set_id"])
        metadata = source.get("project_metadata", {})
        data = source.get("export_data", {})
        original = source.get("original_data", {})
        stage = metadata.get("stage", "")
        if str(stage).isdigit(): stage = int(stage)
        values = [datetime.strptime(zoom["date"], "%Y%m%d"), metadata.get("line", ""), stage,
                  data.get("element", original.get("elemento", "")),
                  data.get("anomaly", original.get("anomalia", "")), zoom["time"], zoom["sequence"],
                  data.get("criticality", original.get("criticidad", ""))]
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str): cell.data_type = "s"
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
        sheet.cell(sheet.max_row, 1).number_format = "dd/mm/yyyy"
        for column in (6, 7): sheet.cell(sheet.max_row, column).number_format = "@"
    for column, width in zip("ABCDEFGH", (14, 28, 10, 30, 52, 22, 24, 16)):
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    workbook.save(output_path)
    workbook.close()
    return output_path
